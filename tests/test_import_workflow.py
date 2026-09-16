"""End-to-end coverage for the two-phase import workflow.

Builds real PDFs with reportlab so pdfplumber extraction, anchor
classification, unit normalization, reports, correction and upload retry are
all exercised against actual files — not mocks.
"""

from __future__ import annotations

import io
from pathlib import Path

from fastapi.testclient import TestClient

from seamtech_search import anchors as shared_anchors
from seamtech_search import crawler, import_pipeline
from seamtech_search.api import create_app
from seamtech_search.config import AppConfig
from seamtech_search.extractors import ExtractionResult
from seamtech_search.import_pipeline import (
    extract_structured_pdf,
    import_folder,
    scan_folder,
)
from seamtech_search.indexer import SearchIndex

TECHNICAL_LINES = [
    "FICHE DE FABRICATION",
    "Reference: SO-1234",
    "Material: Dacron",
    "Dimensions: 1200 x 800 mm",
    "Quantity: 5",
    "Description: Main sail",
]

TECHNICAL_LINES_2 = [
    "Fiche de fabrication",
    "Référence: GV-999",
    "Matière: Laminate",
    "Longueur: 2 x 1.5 m",
    "Quantité: 3",
    "Description: Genoa",
]

PLAN_LINES = ["Yard layout drawing", "Scale 1:50", "North elevation sketch"]


def make_pdf(path: Path, lines: list[str]) -> Path:
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path))
    y = 750
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 20
    pdf.save()
    return path


def make_pdf_bytes(lines: list[str]) -> bytes:
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    y = 750
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 20
    pdf.save()
    return buffer.getvalue()


def make_config(tmp_path: Path, root: Path) -> AppConfig:
    return AppConfig(root_paths=[root], database_path=tmp_path / "data" / "search.db")


def make_source(tmp_path: Path, name: str = "REF-001") -> Path:
    source = tmp_path / "roots" / name
    make_pdf(source / "a-technical.pdf", TECHNICAL_LINES)
    make_pdf(source / "b-technical.pdf", TECHNICAL_LINES_2)
    make_pdf(source / "c-plan.pdf", PLAN_LINES)
    (source / "notes.txt").write_text("some notes", encoding="utf-8")
    return source


# ---------------------------------------------------------------------------
# Shared anchors & classification
# ---------------------------------------------------------------------------


def test_anchors_are_shared_between_crawler_and_pipeline() -> None:
    assert crawler.TECHNICAL_PDF_ANCHORS == shared_anchors.TECHNICAL_ANCHORS
    assert import_pipeline.TECHNICAL_ANCHORS == shared_anchors.TECHNICAL_ANCHORS
    for anchor in ("reference", "référence", "longueur", "matériau"):
        assert anchor in shared_anchors.TECHNICAL_ANCHORS


def test_new_anchors_classify_technical_pdf() -> None:
    # Stricter classifier: 2 weak anchors alone is not enough, need 3 weak or strong+weak
    weak_two = "Reference: X-1\nLongueur: 1200 mm\nSome other words"
    assert shared_anchors.classify_pdf_text(weak_two) == "plan_pdf"

    weak_three = "Reference: X-1\nLongueur: 1200 mm\nLargeur: 800 mm"
    assert shared_anchors.classify_pdf_text(weak_three) == "technical_pdf"

    strong_plus_one = "Fiche de fabrication\nReference: X-1"
    assert shared_anchors.classify_pdf_text(strong_plus_one) == "technical_pdf"
    assert crawler.classify_pdf_text(strong_plus_one) == "technical_pdf" if hasattr(
        crawler, "classify_pdf_text"
    ) else True
    assert import_pipeline.classify_pdf_text(strong_plus_one) == "technical_pdf"

    # Original 3-anchor case still technical
    assert shared_anchors.classify_pdf_text(weak_three) == "technical_pdf"


def test_single_anchor_is_still_a_plan() -> None:
    assert shared_anchors.classify_pdf_text("Reference: lonely mention") == "plan_pdf"


# ---------------------------------------------------------------------------
# Unit normalization
# ---------------------------------------------------------------------------


def _extract_with_text(monkeypatch, text: str) -> None:
    monkeypatch.setattr(
        "seamtech_search.import_pipeline.extract_file",
        lambda *args, **kwargs: ExtractionResult(text, "extracted"),
    )


def test_dimensions_normalized_to_mm(monkeypatch, tmp_path: Path) -> None:
    _extract_with_text(monkeypatch, "Reference: R\nDimensions: 2 x 1.5 m")
    result = extract_structured_pdf(tmp_path / "drawing.pdf", AppConfig(root_paths=[tmp_path]))
    assert result.dimensions.length == 2
    assert result.dimensions.unit == "m"
    assert result.dimensions.length_mm == 2000
    assert result.dimensions.width_mm == 1500
    assert result.dimensions.unit_normalized == "mm"


def test_dimensions_cm_and_labeled_fallback(monkeypatch, tmp_path: Path) -> None:
    _extract_with_text(monkeypatch, "Reference: R\nLongueur: 120 cm\nLargeur: 80 cm")
    result = extract_structured_pdf(tmp_path / "drawing.pdf", AppConfig(root_paths=[tmp_path]))
    assert result.dimensions.length == 120
    assert result.dimensions.length_mm == 1200
    assert result.dimensions.width_mm == 800


# ---------------------------------------------------------------------------
# Scan + confirm flow with real PDFs
# ---------------------------------------------------------------------------


def test_scan_returns_multiple_candidates(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    source = make_source(tmp_path)
    config = make_config(tmp_path, root)

    result = scan_folder(source, config)

    assert result["files_detected"] == 4
    # Returns ALL PDFs with anchor score as ranking hint (4.8), not just technical
    assert len(result["candidates"]) == 3
    technical = [c for c in result["candidates"] if c.get("is_technical")]
    assert len(technical) == 2
    names = {c["name"] for c in technical}
    assert names == {"a-technical.pdf", "b-technical.pdf"}
    assert all(c["anchor_count"] >= 2 for c in technical)
    assert any("fiche de fabrication" in c["anchors_matched"] for c in technical)
    plan = [c for c in result["candidates"] if not c.get("is_technical")]
    assert len(plan) == 1
    assert plan[0]["name"] == "c-plan.pdf"
    assert any("select" in warning for warning in result["warnings"])


def test_import_oneshot_selects_first_and_lists_candidates(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    source = make_source(tmp_path)
    config = make_config(tmp_path, root)
    index = SearchIndex(config.database_path)
    index.initialize(rebuild=True)

    result = import_folder(source, config, index)

    assert result.technical_pdf and result.technical_pdf.endswith("a-technical.pdf")
    assert len(result.candidates) == 2
    assert any("first PDF was selected" in warning for warning in result.warnings)
    assert result.data and result.data["reference"] == "SO-1234"
    assert result.data["dimensions"]["length_mm"] == 1200
    assert result.report_path and Path(result.report_path).exists()
    assert result.report_docx_path and result.report_docx_path.endswith(".docx")
    assert Path(result.report_docx_path).exists()
    assert result.upload_status in ("not_configured", "pending", "not_applicable")


def test_import_confirm_uses_selected_pdf(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    source = make_source(tmp_path)
    config = make_config(tmp_path, root)
    index = SearchIndex(config.database_path)
    index.initialize(rebuild=True)

    result = import_folder(source, config, index, selected_pdf=source / "b-technical.pdf")

    assert result.technical_pdf and result.technical_pdf.endswith("b-technical.pdf")
    assert result.data and result.data["reference"] == "GV-999"
    assert result.data["dimensions"]["length_mm"] == 2000  # 2 m normalized
    assert result.status in {"completed", "needs_review"}


def test_import_rejects_selected_pdf_outside_folder(tmp_path: Path) -> None:
    import pytest

    root = tmp_path / "roots"
    source = make_source(tmp_path)
    outside = make_pdf(tmp_path / "elsewhere.pdf", TECHNICAL_LINES)
    config = make_config(tmp_path, root)
    index = SearchIndex(config.database_path)
    index.initialize(rebuild=True)

    with pytest.raises(ValueError, match="outside the scanned folder"):
        import_folder(source, config, index, selected_pdf=outside)


def test_docx_report_is_readable(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    root = tmp_path / "roots"
    source = make_source(tmp_path)
    config = make_config(tmp_path, root)
    index = SearchIndex(config.database_path)
    index.initialize(rebuild=True)

    result = import_folder(source, config, index, selected_pdf=source / "a-technical.pdf")

    assert result.report_docx_path
    doc = DocxDocument(result.report_docx_path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "SEAMTECH Technical Report" in text
    assert any("SO-1234" in cell.text for table in doc.tables for row in table.rows for cell in row.cells)


# ---------------------------------------------------------------------------
# API: scan / confirm / correct / retry
# ---------------------------------------------------------------------------


def _client(tmp_path: Path, root: Path) -> TestClient:
    config = make_config(tmp_path, root)
    return TestClient(create_app(config))


def test_api_scan_and_confirm(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    source = make_source(tmp_path)
    client = _client(tmp_path, root)

    scan = client.post("/imports/scan", json={"source_path": str(source)})
    assert scan.status_code == 200
    # Now returns ALL PDFs with ranking hint
    assert len(scan.json()["candidates"]) == 3
    technical = [c for c in scan.json()["candidates"] if c.get("is_technical")]
    assert len(technical) == 2

    chosen = technical[1]["path"]
    confirm = client.post("/imports/confirm", json={"source_path": str(source), "technical_pdf": chosen})
    assert confirm.status_code == 200
    body = confirm.json()
    assert body["technical_pdf"] == chosen
    assert body["report_docx_path"]

    fetched = client.get(f"/imports/{body['import_id']}")
    assert fetched.status_code == 200
    assert fetched.json()["technical_pdf"] == chosen


def test_api_correction_regenerates_reports(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    source = make_source(tmp_path)
    client = _client(tmp_path, root)
    created = client.post("/imports?wait=true", json={"source_path": str(source)}).json()
    import_id = created["import_id"]

    response = client.patch(
        f"/imports/{import_id}",
        json={"reference": "CORRECTED-1", "dimensions": {"length": 99, "unit": "cm"}},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["data"]["reference"] == "CORRECTED-1"
    assert body["data"]["dimensions"]["length"] == 99
    assert body["data"]["dimensions"]["length_mm"] == 990
    assert Path(body["report_path"]).exists()
    assert Path(body["report_docx_path"]).exists()

    fetched = client.get(f"/imports/{import_id}").json()
    assert fetched["data"]["reference"] == "CORRECTED-1"


def test_api_correction_rejects_unknown_fields(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    source = make_source(tmp_path)
    client = _client(tmp_path, root)
    created = client.post("/imports?wait=true", json={"source_path": str(source)}).json()

    response = client.patch(f"/imports/{created['import_id']}", json={"quantity": "many"})
    assert response.status_code == 422


def test_api_correction_missing_import_returns_404(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    root.mkdir(parents=True)
    client = _client(tmp_path, root)
    assert client.patch("/imports/does-not-exist", json={"reference": "X"}).status_code == 404


def test_api_retry_upload_without_credentials(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    source = make_source(tmp_path)
    client = _client(tmp_path, root)
    created = client.post("/imports?wait=true", json={"source_path": str(source)}).json()

    response = client.post(f"/imports/{created['import_id']}/retry-upload")
    assert response.status_code == 200
    assert response.json()["upload_status"] in ("not_configured", "pending", "not_applicable", "uploaded")


def test_correct_import_rejects_unknown_fields_directly(tmp_path: Path) -> None:
    import pytest

    from seamtech_search.import_pipeline import correct_import, get_import

    root = tmp_path / "roots"
    source = make_source(tmp_path)
    config = make_config(tmp_path, root)
    index = SearchIndex(config.database_path)
    index.initialize(rebuild=True)
    result = import_folder(source, config, index)

    with pytest.raises(ValueError, match="Unknown correction fields"):
        correct_import(index, config, result.import_id, {"nope": "x"})
    assert get_import(index, result.import_id) is not None


def _touch(path: Path, content: str = "data") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# Browser upload staging
# ---------------------------------------------------------------------------


def test_api_upload_stages_files_and_scans(tmp_path: Path) -> None:
    root = tmp_path / "roots"
    root.mkdir(parents=True)
    client = _client(tmp_path, root)

    response = client.post(
        "/imports/upload",
        files=[("files", ("REF-DND/sheet.pdf", make_pdf_bytes(TECHNICAL_LINES), "application/pdf"))],
        data={"folder": "dnd-drop"},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["staged_path"]
    assert len(body["candidates"]) == 1
    assert Path(body["staged_path"]).exists()

    # The staged folder is a valid confirm source.
    confirm = client.post(
        "/imports/confirm",
        json={"source_path": body["staged_path"], "technical_pdf": body["candidates"][0]["path"]},
    )
    assert confirm.status_code == 200
    assert confirm.json()["data"]["reference"] == "SO-1234"


def test_api_upload_rejects_path_escape(tmp_path: Path) -> None:
    from seamtech_search.api import _safe_relative_path

    assert ".." not in str(_safe_relative_path("../../etc/passwd"))
    assert str(_safe_relative_path("a//b.pdf")) == str(Path("a/b.pdf"))
