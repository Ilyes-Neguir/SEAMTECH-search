from pathlib import Path

import openpyxl

from seamtech_search.config import AppConfig
from seamtech_search.extractors import ExtractionResult
from seamtech_search.import_pipeline import (
    extract_excel_summary,
    extract_structured_pdf,
    generate_docx_report,
    generate_report,
    import_folder,
    scan_folder,
)
from seamtech_search.indexer import SearchIndex


def _extract_with_text(monkeypatch, text: str) -> None:
    monkeypatch.setattr(
        "seamtech_search.import_pipeline.extract_file",
        lambda *args, **kwargs: ExtractionResult(text, "extracted"),
    )


def test_description_extracts_value_on_same_line(monkeypatch, tmp_path: Path) -> None:
    _extract_with_text(monkeypatch, "Fiche de fabrication 7792-SO - Porte coulissante")

    result = extract_structured_pdf(tmp_path / "drawing.pdf", AppConfig(root_paths=[tmp_path]))

    assert result.description == "7792-SO - Porte coulissante"


def test_description_header_does_not_capture_following_line(monkeypatch, tmp_path: Path) -> None:
    _extract_with_text(
        monkeypatch,
        "FICHE DE FABRICATION\nCommande 7792-SO\nReference: 7792-SO-GV",
    )

    result = extract_structured_pdf(tmp_path / "drawing.pdf", AppConfig(root_paths=[tmp_path]))

    assert result.description is None
    assert result.reference == "7792-SO-GV"


def test_extract_excel_summary(tmp_path: Path) -> None:
    wb_path = tmp_path / "test_bom.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nomenclature"
    ws.append(["Repere", "Designation", "Matiere", "Quantite"])
    ws.append(["01", "Panneau central", "Dacron 380g", "1"])
    ws.append(["02", "Renfort point d'ecoute", "Spectra 450g", "2"])
    ws.append(["03", "Ralingue guindant", "Cordage Polyester 10mm", "12m"])
    wb.save(wb_path)

    summary = extract_excel_summary(wb_path)
    assert summary is not None
    assert summary.filename == "test_bom.xlsx"
    assert summary.total_sheets == 1
    assert "Nomenclature" in summary.sheet_names
    sheet_sum = summary.sheets[0]
    assert sheet_sum.sheet_name == "Nomenclature"
    assert sheet_sum.row_count == 4
    assert sheet_sum.column_count == 4
    assert sheet_sum.headers == ["Repere", "Designation", "Matiere", "Quantite"]
    assert len(sheet_sum.sample_rows) == 3
    assert sheet_sum.sample_rows[0][1] == "Panneau central"


def test_scan_folder_detects_excel(tmp_path: Path) -> None:
    config = AppConfig(root_paths=[tmp_path], min_free_bytes=0)
    (tmp_path / "doc.pdf").write_text("dummy")
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "bom.xlsx")

    scan_res = scan_folder(tmp_path, config)
    assert scan_res["files_detected"] == 2
    assert len(scan_res["excel_candidates"]) == 1
    assert scan_res["excel_candidates"][0]["name"] == "bom.xlsx"


def test_generate_reports_with_excel(monkeypatch, tmp_path: Path) -> None:
    test_text = (
        "FICHE DE FABRICATION\nReference: REF-999\nMatiere: Kevlar\n"
        "Quantite: 10\nDimensions: 500x300 mm\nDescription: Piece test"
    )
    _extract_with_text(monkeypatch, test_text)
    data = extract_structured_pdf(tmp_path / "tech.pdf", AppConfig(root_paths=[tmp_path]))

    wb_path = tmp_path / "items.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["Item", "Qty", "Spec"])
    ws.append(["Sleeve", "4", "PVC"])
    wb.save(wb_path)

    excel_summary = extract_excel_summary(wb_path)

    pdf_out = tmp_path / "report.pdf"
    docx_out = tmp_path / "report.docx"

    generate_report(
        data=data,
        output_path=pdf_out,
        excel_summary=excel_summary,
    )
    assert pdf_out.exists()
    assert pdf_out.stat().st_size > 0

    generate_docx_report(
        data=data,
        output_path=docx_out,
        excel_summary=excel_summary,
    )
    assert docx_out.exists()
    assert docx_out.stat().st_size > 0


def test_end_to_end_import_with_excel(tmp_path: Path) -> None:
    source_folder = tmp_path / "reference_job"
    source_folder.mkdir()

    shutil_fixture_pdf = Path("sample_data/CLIENT-123/fiche-technique.pdf")
    shutil_fixture_xlsx = Path("sample_data/CLIENT-123/nomenclature-voile.xlsx")

    (source_folder / "fiche-technique.pdf").write_bytes(shutil_fixture_pdf.read_bytes())
    (source_folder / "nomenclature-voile.xlsx").write_bytes(shutil_fixture_xlsx.read_bytes())

    config = AppConfig(
        root_paths=[tmp_path],
        database_path=tmp_path / "search.db",
        min_free_bytes=0,
    )
    index = SearchIndex(config.database_path)
    index.initialize(rebuild=True)

    result = import_folder(source_folder, config, index)

    assert result.status == "completed"
    assert result.technical_pdf is not None
    assert result.excel_file is not None
    assert "nomenclature-voile.xlsx" in result.excel_file
    assert result.excel_summary is not None
    assert result.excel_summary["total_sheets"] == 2
    assert result.report_path is not None
    assert result.report_docx_path is not None
