from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from seamtech_search.api import create_app
from seamtech_search.config import AppConfig
from seamtech_search.indexer import SearchIndex
from seamtech_search.models import Document


def test_configured_token_protects_search(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    config = AppConfig(root_paths=[root], database_path=tmp_path / "search.db", auth_token="secret")
    client = TestClient(create_app(config))

    assert client.get("/search?q=client").status_code == 401
    assert client.get("/search?q=client", headers={"X-SEAMTECH-TOKEN": "secret"}).status_code == 200
    assert client.get("/health").status_code == 401
    assert client.get("/health", headers={"X-SEAMTECH-TOKEN": "secret"}).status_code == 200


def test_search_pagination_contract(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    database = tmp_path / "search.db"
    index = SearchIndex(database)
    index.initialize(rebuild=True)
    for number in range(3):
        path = root / f"reference-{number}.txt"
        index.upsert_document(Document(path, path.name, root, ".txt", 10, float(number), False, "factory reference"))
    config = AppConfig(root_paths=[root], database_path=database)
    client = TestClient(create_app(config))
    response = client.get("/search?q=reference&limit=2&offset=1")
    assert response.status_code == 200
    payload = response.json()
    assert payload["offset"] == 1
    assert payload["limit"] == 2
    assert payload["count"] == 2
    assert payload["has_more"] is False


def test_local_mode_allows_search_without_token(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    config = AppConfig(root_paths=[root], database_path=tmp_path / "search.db")
    client = TestClient(create_app(config))
    assert client.get("/search?q=client").status_code == 200


def test_search_exposes_extraction_status(tmp_path: Path) -> None:
    root = tmp_path / "root"
    root.mkdir()
    database = tmp_path / "search.db"
    index = SearchIndex(database)
    index.initialize(rebuild=True)
    path = root / "drawing.dwg"
    index.upsert_document(
        Document(
            path,
            path.name,
            root,
            ".dwg",
            10,
            1.0,
            False,
            extraction_status="unavailable",
            extraction_detail="unsupported file type",
        )
    )
    client = TestClient(create_app(AppConfig(root_paths=[root], database_path=database)))

    result = client.get("/search?q=drawing").json()["results"][0]

    assert result["extraction_status"] == "unavailable"
    assert result["extraction_detail"] == "unsupported file type"


def test_api_import_with_excel_support(tmp_path: Path) -> None:
    folder = tmp_path / "job_folder"
    folder.mkdir()

    # Copy sample fixture files
    shutil_fixture_pdf = Path("sample_data/CLIENT-123/fiche-technique.pdf")
    (folder / "fiche-technique.pdf").write_bytes(shutil_fixture_pdf.read_bytes())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["Ref", "Desc", "Matiere", "Qty"])
    ws.append(["001", "Bout d'écoute", "Polyester", "2"])
    wb.save(folder / "bom.xlsx")

    config = AppConfig(
        root_paths=[tmp_path],
        database_path=tmp_path / "search.db",
        min_free_bytes=0,
    )
    client = TestClient(create_app(config))

    # Test Scan endpoint
    scan_resp = client.post("/imports/scan", json={"source_path": str(folder)})
    assert scan_resp.status_code == 200
    scan_data = scan_resp.json()
    assert len(scan_data["candidates"]) == 1
    assert len(scan_data["excel_candidates"]) == 1
    assert scan_data["excel_candidates"][0]["name"] == "bom.xlsx"

    # Test Confirm endpoint with excel_file
    confirm_resp = client.post(
        "/imports/confirm?wait=true",
        json={
            "source_path": str(folder),
            "technical_pdf": str(folder / "fiche-technique.pdf"),
            "excel_file": str(folder / "bom.xlsx"),
        },
    )
    assert confirm_resp.status_code == 200
    confirm_data = confirm_resp.json()
    assert confirm_data["status"] == "completed"
    assert confirm_data["excel_file"] is not None
    assert "bom.xlsx" in confirm_data["excel_file"]
    assert confirm_data["excel_summary"] is not None
    assert confirm_data["excel_summary"]["total_sheets"] == 1
